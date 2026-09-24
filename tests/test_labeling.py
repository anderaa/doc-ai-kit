"""The labeling step: the sample, the sheet, and reading a person's answers back."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import pytest
import yaml
from click.testing import CliRunner
from conftest import document_for, scripted_lm
from openpyxl import load_workbook

from doc_ai_kit.cli import Project, cli
from doc_ai_kit.dataset import LabelRecord, load_labels
from doc_ai_kit.extract import doc_id_from_name
from doc_ai_kit.labeling import (
    CellError,
    LabelingError,
    LabelPlan,
    SheetRow,
    draw_sample,
    import_rows,
    parse_cell,
    read_sheet,
    render_cell,
    write_sheet,
)
from doc_ai_kit.metric import build_metric
from doc_ai_kit.registry import Registry
from doc_ai_kit.scaffold_writer import ScaffoldOptions, create_project
from doc_ai_kit.splits import make_splits
from doc_ai_kit.state import derive

DOCUMENT = (
    "MASTER SERVICES AGREEMENT No. 00417\n"
    "This Agreement shall be governed by the laws of\nthe State of Texas.\n"
    "Signed by Jane Q. Smith and Omar Haddad."
)
GOVERNING = "This Agreement shall be governed by the laws of the State of Texas."


@pytest.fixture(scope="module")
def registry(fixtures_dir: Path) -> Registry:
    return Registry.from_yaml(fixtures_dir / "tasks_all_types.yaml")


def _parse(registry: Registry, task_id: str, text: str, document: str = DOCUMENT) -> Any:
    value, _warnings = parse_cell(registry.by_id(task_id), text, document)
    return value


# the sample


def test_the_sample_is_fixed_by_its_seed() -> None:
    corpus = [f"d{index:03d}" for index in range(200)]
    first = draw_sample(corpus, 40, 0.3, seed=7)
    assert first.sampled == draw_sample(list(reversed(corpus)), 40, 0.3, seed=7).sampled
    assert first.holdout == draw_sample(corpus, 40, 0.3, seed=7).holdout
    assert first.sampled != draw_sample(corpus, 40, 0.3, seed=8).sampled
    assert len(first.sampled) == 40 and len(first.holdout) == 12
    assert set(first.holdout) <= set(first.sampled)
    assert first.inclusion_probability == pytest.approx(0.2)


def test_the_sample_is_capped_at_the_corpus() -> None:
    plan = draw_sample(["a", "b", "c"], 10, 0.3, seed=1)
    assert plan.sampled == ["a", "b", "c"] and plan.inclusion_probability == 1.0


def test_a_plan_refuses_a_holdout_that_was_shown_model_answers() -> None:
    with pytest.raises(LabelingError, match="no longer blind"):
        LabelPlan(seed=1, corpus_size=3, sampled=["a", "b"], holdout=["b"], prefilled=["a", "b"])


# reading cells


@pytest.mark.parametrize("text,expected", [("yes", True), ("No", False), ("Y", True), ("", None), ("N/A", None)])
def test_binary_cells(registry: Registry, text: str, expected: Any) -> None:
    assert _parse(registry, "has_arbitration_clause", text) is expected


def test_enum_cells_take_synonyms_and_reject_anything_else(registry: Registry) -> None:
    assert _parse(registry, "filing_state", "California") == "CA"
    assert _parse(registry, "filing_state", "tx") == "TX"
    assert _parse(registry, "filing_state", "") is None
    with pytest.raises(CellError, match="not one of"):
        _parse(registry, "filing_state", "Ohio")


def test_multilabel_cells_split_on_semicolons(registry: Registry) -> None:
    assert _parse(registry, "covered_products", "Support; hardware;") == ["hardware", "support"]
    assert _parse(registry, "covered_products", "") == []
    with pytest.raises(CellError, match="'firmware' is not one of"):
        _parse(registry, "covered_products", "hardware; firmware")


def test_text_cells_keep_what_was_typed(registry: Registry) -> None:
    """A leading zero is part of a contract number; Excel dropping it is why the sheet is text."""
    assert _parse(registry, "contract_number", "00417") == "00417"
    assert _parse(registry, "contract_number", "n/a") is None
    assert _parse(registry, "signatories", "Jane Q. Smith; Omar Haddad") == ["Jane Q. Smith", "Omar Haddad"]


def test_wrapping_quote_marks_never_reach_a_label(registry: Registry) -> None:
    """Seen live: the model answered '"MSA-2024-1014"', quote marks included."""
    assert render_cell(registry.by_id("contract_number"), '"MSA-2024-1014"') == "MSA-2024-1014"
    assert _parse(registry, "contract_number", "\u201cMSA-2024-1014\u201d") == "MSA-2024-1014"
    assert _parse(registry, "signatories", '"Ana Ruiz"; Omar Haddad') == ["Ana Ruiz", "Omar Haddad"]


def test_numeric_cells_accept_the_ways_people_write_money(registry: Registry) -> None:
    for text in ("$1.25M", "1,250,000 USD", "1250000"):
        assert _parse(registry, "contract_value", text) == {"value": 1250000.0, "unit": "USD"}, text
    with pytest.raises(CellError, match="not a number"):
        _parse(registry, "contract_value", "about a million")


def test_numeric_cells_warn_on_a_different_unit(registry: Registry) -> None:
    value, warnings = parse_cell(registry.by_id("contract_value"), "€40,000", DOCUMENT)
    assert value == {"value": 40000.0, "unit": "EUR"}
    assert warnings and "EUR" in warnings[0]


def test_date_cells_keep_their_precision(registry: Registry) -> None:
    assert _parse(registry, "effective_date", "June 15, 2024") == "2024-06-15"
    assert _parse(registry, "effective_date", "2024-06") == "2024-06"
    with pytest.raises(CellError, match="not a date"):
        _parse(registry, "effective_date", "next spring")


def test_a_pasted_passage_becomes_offsets(registry: Registry) -> None:
    """Copied from a PDF viewer, the line breaks differ from the text cache; it still lands."""
    span = _parse(registry, "governing_law_span", GOVERNING)
    assert DOCUMENT[span["start"] : span["end"]].replace("\n", " ") == GOVERNING


def test_a_passage_not_in_the_document_is_refused(registry: Registry) -> None:
    with pytest.raises(CellError, match="not in the extracted text"):
        _parse(registry, "governing_law_span", "This Agreement is governed by the laws of Delaware.")


def test_a_passage_found_twice_is_flagged(registry: Registry) -> None:
    _value, warnings = parse_cell(registry.by_id("governing_law_span"), "Texas", DOCUMENT + " Texas.")
    assert warnings and "more than once" in warnings[0]


def test_a_non_nullable_task_may_not_be_blank(tmp_path: Path) -> None:
    registry = Registry.from_mapping(
        {"tasks": [{"id": "number", "type": "extract_exact", "question": "The number.", "nullable": False}]}
    )
    with pytest.raises(CellError, match="may not be blank"):
        parse_cell(registry.by_id("number"), "", DOCUMENT)


def test_every_label_survives_the_round_trip_through_a_cell(registry: Registry) -> None:
    """What label-sheet writes, import-labels must read back as the same label."""
    span_start = DOCUMENT.index("This Agreement")
    gold = {
        "has_arbitration_clause": False,
        "filing_state": "TX",
        "covered_products": ["hardware", "support"],
        "contract_number": "00417",
        "counterparty": "Contoso Manufacturing LLC",
        "signatories": ["Jane Q. Smith", "Omar Haddad"],
        "contract_value": {"value": 1250000.0, "unit": "USD"},
        "effective_date": "2024-06-15",
        "governing_law_span": {"start": span_start, "end": DOCUMENT.index("Texas.") + len("Texas.")},
    }
    metric = build_metric(registry)
    for task in registry:
        cell = render_cell(task, gold[task.id], DOCUMENT)
        value, _warnings = parse_cell(task, cell, DOCUMENT)
        assert value == gold[task.id], (task.id, cell)
        assert metric.score_task(task, gold, {**gold, task.id: value}).correct, task.id


def test_model_answers_render_as_a_person_would_write_them(registry: Registry) -> None:
    from doc_ai_kit.values import PartialDate, Quantity, Span

    assert render_cell(registry.by_id("contract_value"), Quantity(value=3250000.0, unit="USD")) == "3,250,000 USD"
    assert render_cell(registry.by_id("contract_value"), Quantity(value=0.5, unit=None)) == "0.5"
    assert render_cell(registry.by_id("effective_date"), PartialDate(value="2024-06", granularity="month")) == "2024-06"
    assert render_cell(registry.by_id("governing_law_span"), Span(start=1, end=4, text="abc")) == "abc"
    assert render_cell(registry.by_id("has_arbitration_clause"), True) == "yes"
    assert render_cell(registry.by_id("filing_state"), None) == ""


# the sheet


@pytest.fixture
def plan() -> LabelPlan:
    return LabelPlan(seed=1, corpus_size=10, sampled=["a", "b", "c", "d"], holdout=["d"], prefilled=["a", "b"])


def _toy(fixtures_dir: Path) -> Registry:
    return Registry.from_yaml(fixtures_dir / "toy_tasks.yaml")


def test_the_sheet_is_one_tab_of_file_names_and_labels(tmp_path: Path, fixtures_dir: Path, plan: LabelPlan) -> None:
    registry = _toy(fixtures_dir)
    answers = {doc_id: {"flag": True, "state": "CA", "number": "0042"} for doc_id in ("a", "b", "d")}
    path = tmp_path / "labels.xlsx"
    write_sheet(path, registry, plan, answers, {"b": "odd scan"}, tmp_path)

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["labels"]
    sheet = workbook["labels"]
    rows = {doc_id_from_name(row[0]): row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert [cell.value for cell in sheet[1]] == ["file_name", "flag", "state", "number", "notes"]
    assert sheet["A2"].value == "a.pdf"
    assert rows["a"][1:] == ("yes", "CA", "0042", None)
    assert rows["b"][4] == "odd scan"
    assert rows["c"][1:4] == (None, None, None)
    # a holdout row shows only what it is given; the prefill step is what never runs the model on it
    assert rows["d"][1:4] == ("yes", "CA", "0042")
    shaded = {doc_id_from_name(row[0].value) for row in sheet.iter_rows(min_row=2) if row[0].fill.fill_type}
    assert shaded == {"c", "d"}, "rows to label from scratch are the ones not prefilled, holdout included"
    assert sheet.cell(row=2, column=4).number_format == "@"


def test_read_sheet_reads_xlsx_and_csv_alike(tmp_path: Path, fixtures_dir: Path, plan: LabelPlan) -> None:
    registry = _toy(fixtures_dir)
    path = tmp_path / "labels.xlsx"
    write_sheet(path, registry, plan, {"a": {"flag": True, "state": "CA", "number": "0042"}}, {}, tmp_path)
    from_xlsx = read_sheet(path, registry)

    csv_path = tmp_path / "labels.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["doc_id", "flag", "state", "number", "notes"])
        writer.writerow(["a", "yes", "CA", "0042", ""])
    from_csv = read_sheet(csv_path, registry)
    assert from_xlsx[0].cells == from_csv[0].cells == {"flag": "yes", "state": "CA", "number": "0042"}


def test_a_renamed_tab_and_extra_columns_are_fine(tmp_path: Path, fixtures_dir: Path, plan: LabelPlan) -> None:
    registry = _toy(fixtures_dir)
    path = tmp_path / "labels.xlsx"
    write_sheet(path, registry, plan, {"a": {"flag": True, "state": "CA", "number": "1"}}, {}, tmp_path)
    workbook = load_workbook(path)
    workbook["labels"].title = "My labels"
    workbook["My labels"].cell(row=1, column=7).value = "my own column"
    workbook.save(path)
    assert read_sheet(path, registry)[0].cells["state"] == "CA"


def test_read_sheet_names_a_missing_task_column(tmp_path: Path, fixtures_dir: Path) -> None:
    path = tmp_path / "labels.csv"
    path.write_text("file_name,flag\na.pdf,yes\n", encoding="utf-8")
    with pytest.raises(LabelingError, match="state, number"):
        read_sheet(path, _toy(fixtures_dir))


# importing


def _row(doc_id: str, notes: str = "", **cells: str) -> SheetRow:
    return SheetRow(doc_id=doc_id, cells={"flag": "yes", "state": "CA", "number": "7", **cells}, notes=notes)


def _texts(plan: LabelPlan) -> dict[str, str]:
    return {doc_id: document_for(doc_id) for doc_id in plan.sampled}


def test_import_records_how_each_document_was_labeled(fixtures_dir: Path, plan: LabelPlan) -> None:
    rows = [_row("a"), _row("b"), _row("c", notes="skip: not a contract"), _row("d")]
    result = import_rows(_toy(fixtures_dir), plan, rows, _texts(plan))
    assert result.ok, result.problems
    modes = {record.doc_id: record.labeling_mode for record in result.records}
    assert modes == {"a": "corrected", "b": "corrected", "d": "blind"}
    assert result.skipped == {"c": "not a contract"}
    assert all(record.inclusion_probability == pytest.approx(0.4) for record in result.records)
    assert result.records[0].labels == {"flag": True, "state": "CA", "number": "7"}


def test_an_empty_row_means_unfinished(fixtures_dir: Path, plan: LabelPlan) -> None:
    """A blank cell can mean "no answer", but a row with nothing at all was never labeled."""
    empty = SheetRow(doc_id="c", cells={"flag": "", "state": "", "number": ""})
    result = import_rows(_toy(fixtures_dir), plan, [_row("a"), _row("b"), empty, _row("d")], _texts(plan))
    assert not result.ok and any("c: nothing filled in yet" in problem for problem in result.problems)

    considered = SheetRow(doc_id="c", cells={"flag": "", "state": "", "number": ""}, notes="answers none of these")
    result = import_rows(_toy(fixtures_dir), plan, [_row("a"), _row("b"), considered, _row("d")], _texts(plan))
    assert result.ok, result.problems
    assert next(r for r in result.records if r.doc_id == "c").labels == {"flag": None, "state": None, "number": None}


@pytest.mark.parametrize("note", ["skip: unreadable", "Skip - unreadable", "SKIP unreadable", "skip \u2013 unreadable"])
def test_skip_notes_are_read_generously(note: str) -> None:
    assert SheetRow(doc_id="x", cells={}, notes=note).skip_reason == "unreadable"


def test_a_note_mentioning_skip_later_is_not_a_skip() -> None:
    assert SheetRow(doc_id="x", cells={}, notes="checked; did not skip anything").skip_reason is None
    assert SheetRow(doc_id="x", cells={}, notes="skipping clause 4 was deliberate").skip_reason is None


def test_import_lists_every_problem_at_once(fixtures_dir: Path, plan: LabelPlan) -> None:
    rows = [
        _row("a", state="Ohio", flag="maybe"),
        _row("c", notes="skip"),
        _row("zz"),
        _row("d"),
        _row("d"),
    ]
    result = import_rows(_toy(fixtures_dir), plan, rows, _texts(plan))
    assert not result.ok
    joined = "\n".join(result.problems)
    for fragment in (
        "a / state",
        "a / flag",
        "c: skipped with no reason",
        "zz: not in the labeling sample",
        "d: appears in more than one row",
        "1 sampled document(s) have no row: b",
    ):
        assert fragment in joined, fragment
    assert not any(record.doc_id == "a" for record in result.records)


# splits


def test_make_splits_keeps_a_holdout_drawn_before_labeling(fixtures_dir: Path) -> None:
    registry = _toy(fixtures_dir)
    records = [
        LabelRecord(doc_id=f"d{index:02d}", labels={"flag": index % 2 == 0, "state": ["CA", "NY", "TX"][index % 3]})
        for index in range(60)
    ]
    holdout = [f"d{index:02d}" for index in range(0, 60, 4)]
    splits = make_splits(registry, records, seed=3, holdout=holdout)
    assert splits.holdout == sorted(holdout)
    assert sorted(splits.train + splits.val + splits.holdout) == sorted(record.doc_id for record in records)
    assert "holdout drawn before labeling" in splits.strategy


def test_make_splits_refuses_an_unlabeled_holdout_document(fixtures_dir: Path) -> None:
    records = [LabelRecord(doc_id=f"d{index}", labels={"flag": True, "state": "CA"}) for index in range(10)]
    with pytest.raises(ValueError, match="not labeled: missing"):
        make_splits(_toy(fixtures_dir), records, seed=1, holdout=["d1", "missing"])


# the commands, end to end


@pytest.fixture
def project(tmp_path: Path, fixtures_dir: Path) -> Path:
    target = tmp_path / "acme"
    create_project(target, ScaffoldOptions(project_name="Acme"))
    (target / "tasks.yaml").write_text((fixtures_dir / "toy_tasks.yaml").read_text(encoding="utf-8"))
    config = yaml.safe_load((target / "config.yaml").read_text(encoding="utf-8"))
    config["models"]["task"] = config["models"]["reflection"] = "fake/scripted"
    config["splits"].update(support_floor=1, measurable_floor=1)
    (target / "config.yaml").write_text(yaml.safe_dump(config), encoding="utf-8")
    text_dir = target / "data" / "text"
    text_dir.mkdir(parents=True, exist_ok=True)
    for index in range(30):
        doc_id = f"doc{index:02d}"
        (text_dir / f"{doc_id}.md").write_text(document_for(doc_id), encoding="utf-8")
    return target


def _invoke(project: Path, *args: str) -> Any:
    result = CliRunner().invoke(cli, ["--project", str(project), *args], catch_exceptions=False)
    return result


def _fill_sheet(path: Path, answer: dict[str, str], only: set[str] | None = None) -> None:
    """Play the labeler: fill in the empty cells."""
    workbook = load_workbook(path)
    sheet = workbook["labels"]
    headers = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2):
        if only is not None and doc_id_from_name(row[0].value) not in only:
            continue
        for column, name in enumerate(headers):
            if name in answer and not row[column].value:
                row[column].value = answer[name]
    workbook.save(path)


def test_labeling_end_to_end(project: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import doc_ai_kit.produce as produce_module

    seen: list[str] = []
    real_produce = produce_module.produce

    def spying_produce(registry: Any, config: Any, program: Any, texts: Any, *args: Any, **kwargs: Any) -> Any:
        seen.extend(texts)
        return real_produce(registry, config, program, texts, *args, **kwargs)

    monkeypatch.setattr(produce_module, "produce", spying_produce)
    monkeypatch.setattr(Project, "configure_lm", lambda self: None)

    result = _invoke(project, "sample-labels", "--count", "20")
    assert result.exit_code == 0, result.output
    plan = json.loads((project / "data" / "label_plan.json").read_text())
    assert len(plan["sampled"]) == 20 and len(plan["holdout"]) == 6
    assert "labeling sample" in (project / "decisions.md").read_text()

    answers = {
        doc_id: {"flag": ("true", "false")[index % 2], "state": ("CA", "NY", "TX")[index % 3], "number": "A-1"}
        for index, doc_id in enumerate(plan["sampled"])
    }
    # the labeler has to choose: prefilled or not
    refused = _invoke(project, "label-sheet")
    assert refused.exit_code != 0 and "--prefill or --no-prefill" in refused.output
    with scripted_lm(answers):
        result = _invoke(project, "label-sheet", "--prefill")
    assert result.exit_code == 0, result.output
    assert set(seen) == set(plan["sampled"]) - set(plan["holdout"]), "the model saw a holdout document"
    plan = json.loads((project / "data" / "label_plan.json").read_text())
    assert sorted(plan["prefilled"]) == sorted(set(plan["sampled"]) - set(plan["holdout"]))

    sheet_path = project / "data" / "labels.xlsx"
    refused = _invoke(project, "label-sheet", "--prefill")
    assert refused.exit_code != 0 and "import-labels" in refused.output

    # an unfinished sheet is refused whole: the holdout rows are still empty
    to_correct = set(plan["sampled"]) - set(plan["holdout"])
    _fill_sheet(sheet_path, {"flag": "no", "state": "CA", "number": "B-2"}, only=to_correct)
    assert any("changed since the last import" in warning for warning in derive(project).warnings)
    result = _invoke(project, "import-labels")
    assert result.exit_code == 1
    assert result.output.count("nothing filled in yet") == 6
    assert not (project / "data" / "labels.jsonl").exists()

    # finished: the holdout labeled from scratch
    _fill_sheet(sheet_path, {"flag": "yes", "state": "TX", "number": "C-3"})
    result = _invoke(project, "import-labels")
    assert result.exit_code == 0, result.output
    assert "Holdout: 6 of 6 labeled" in result.output
    records = load_labels(project / "data" / "labels.jsonl")
    blind = {record.doc_id for record in records if record.labeling_mode == "blind"}
    assert blind == set(plan["holdout"])
    assert {record.labels["state"] for record in records if record.doc_id in blind} == {"TX"}
    # prefilled answers the labeler left alone are kept, not overwritten by the fill
    corrected = [record for record in records if record.doc_id not in blind]
    assert all(record.labels["state"] == answers[record.doc_id]["state"] for record in corrected)

    state = derive(project)
    assert "20 of 20 labeled (6 of 6 holdout)" in state.to_text()
    assert not any("changed since the last import" in warning for warning in state.warnings)

    # rewriting the sheet brings back what was imported, with no model call and no choice to make
    seen.clear()
    assert _invoke(project, "label-sheet", "--force").exit_code == 0
    assert not seen
    rewritten = {
        doc_id_from_name(row[0]): row
        for row in load_workbook(sheet_path)["labels"].iter_rows(min_row=2, values_only=True)
    }
    assert all(rewritten[doc_id][2] == "TX" for doc_id in blind)

    (project / "data" / "annotation_rules.md").write_text("rules", encoding="utf-8")
    result = _invoke(project, "make-splits", "--non-interactive")
    assert result.exit_code == 0, result.output
    splits = json.loads((project / "data" / "splits.json").read_text())
    assert sorted(splits["assignments"]["holdout"]) == sorted(plan["holdout"])
    assert not derive(project).blockers


def test_sample_labels_refuses_to_redraw_once_labeled(project: Path) -> None:
    assert _invoke(project, "sample-labels", "--count", "10").exit_code == 0
    refused = _invoke(project, "sample-labels", "--count", "10")
    assert refused.exit_code != 0
    (project / "data" / "labels.jsonl").write_text("", encoding="utf-8")
    refused = _invoke(project, "sample-labels", "--count", "10", "--force")
    assert refused.exit_code != 0


def test_import_labels_writes_nothing_when_a_cell_is_wrong(project: Path) -> None:
    assert _invoke(project, "sample-labels", "--count", "10").exit_code == 0
    assert _invoke(project, "label-sheet", "--no-prefill").exit_code == 0
    _fill_sheet(project / "data" / "labels.xlsx", {"flag": "yes", "state": "Ohio", "number": "1"})
    result = _invoke(project, "import-labels")
    assert result.exit_code == 1
    assert "nothing was written" in result.output and "state" in result.output
    assert not (project / "data" / "labels.jsonl").exists()


def test_without_prefill_every_row_is_blind(project: Path) -> None:
    assert _invoke(project, "sample-labels", "--count", "10").exit_code == 0
    assert _invoke(project, "label-sheet", "--no-prefill").exit_code == 0
    _fill_sheet(project / "data" / "labels.xlsx", {"flag": "yes", "state": "CA", "number": "1"})
    assert _invoke(project, "import-labels").exit_code == 0
    assert {record.labeling_mode for record in load_labels(project / "data" / "labels.jsonl")} == {"blind"}


def test_rows_are_named_by_the_real_file_name(tmp_path: Path, fixtures_dir: Path, plan: LabelPlan) -> None:
    """Seen in CUAD: capital .PDF, and a space before it that a spreadsheet cell drops."""
    registry = _toy(fixtures_dir)
    names = {"a": "Sonos, Inc. - Manufacturing Agreement .PDF"}
    single = LabelPlan(seed=1, corpus_size=1, sampled=["a"], holdout=[])
    path = tmp_path / "labels.xlsx"
    write_sheet(path, registry, single, {}, {}, tmp_path, file_names=names)
    assert load_workbook(path)["labels"]["A2"].value == names["a"]

    workbook = load_workbook(path)
    workbook["labels"]["A2"].value = "  Sonos, Inc. - Manufacturing Agreement .PDF "
    workbook.save(path)
    assert [row.doc_id for row in read_sheet(path, registry)] == ["Sonos, Inc. - Manufacturing Agreement"]


@pytest.mark.parametrize(
    "cell,warns",
    [
        ("Stryker Corporation and Conformis Inc", True),
        ("Jane Q. Smith and Omar Haddad", True),
        ("Contoso LLC / Fabrikam Industries, Inc.", True),
        ("Jane Q. Smith; Omar Haddad", False),
        ("Johnson & Johnson", False),
        ("Procter and Gamble", False),
        ("Aduro Biotech, Inc.", False),
        ("Smith & Hartley Ltd.", False),
    ],
)
def test_two_entries_in_one_cell_are_flagged(registry: Registry, cell: str, warns: bool) -> None:
    """Reported from a real run: two names in one cell capped that task at 0.92, silently.

    The warning has to stay quiet on ordinary names, or it teaches people to ignore warnings.
    """
    _value, warnings = parse_cell(registry.by_id("signatories"), cell, DOCUMENT)
    assert bool(warnings) is warns, warnings
