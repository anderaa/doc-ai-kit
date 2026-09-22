"""The labeling step: choose documents, hand a person a spreadsheet, read their answers back.

Three things make this harder than writing rows to a file.

The holdout has to be labeled blind, while everything else is labeled faster by correcting
the model's answers. So the holdout is drawn **before** any model runs, at random, and its
rows go out empty. It is drawn at random rather than stratified because nothing is labeled
yet to stratify on -- and a random holdout is also the one whose number is an honest estimate
for the corpus.

People label in their own formats: ``$1.25M``, ``June 15, 2024``, ``TX``. Each cell is read
with the same normalizers used for scoring, then stored in one canonical form, so a
label that could not be scored is caught at import rather than halfway through a run.

Nobody can type character offsets. A span is labeled by pasting the passage, and the
harness finds where it is in the extracted text, as it does for the model's own quotes.

The sheet is ``.xlsx`` rather than CSV because Excel rewrites CSV cells on open: contract
numbers lose their leading zeros and dates turn into its own format. Every cell here is
stored as text, and enum columns get dropdowns.
"""

from __future__ import annotations

import csv
import json
import logging
import random
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from doc_harness.dataset import DatasetError, LabelRecord
from doc_harness.extract import doc_id_from_name
from doc_harness.normalize import LEGAL_SUFFIXES, boolean, is_null, strip_wrapping_quotes
from doc_harness.normalize import date as normalize_date
from doc_harness.normalize import enum as normalize_enum
from doc_harness.normalize import numeric as normalize_numeric
from doc_harness.registry import Registry, TaskType
from doc_harness.spans import locate_quote
from doc_harness.values import PartialDate, Quantity, Span

logger = logging.getLogger(__name__)

PLAN_FILE = "label_plan.json"
SHEET_FILE = "labels.xlsx"
SHEET_NAME = "labels"

# rows are named by the PDF's file name, as the labeler sees it in the folder; a sheet
# naming them by document id instead is read too
FILE_COLUMN = "file_name"
DOC_COLUMN = "doc_id"
NOTES_COLUMN = "notes"

# a note starting "skip: <reason>" sets a document aside, e.g. one that is not a contract at all
_SKIP_NOTE = re.compile(r"^\s*skip\b\s*[:\-\u2013\u2014]?\s*(.*)$", re.IGNORECASE | re.DOTALL)

# shading on the file_name cell of rows to label from scratch: the holdout, and any row the
# model gave no answers for
BLIND_FILL = "FFF2CC"

# several answers in one cell; semicolons, because names and product lines carry commas
LIST_SEPARATOR = ";"

# how to fill in each kind of column, shown on the guide sheet and in header comments
FORMAT_HINTS: dict[TaskType, str] = {
    TaskType.BINARY: "yes or no. Blank if the document does not say.",
    TaskType.MULTICLASS: "One of the allowed values. Blank if none applies.",
    TaskType.MULTILABEL: "Allowed values separated by semicolons, e.g. hardware; support. Blank for none.",
    TaskType.EXTRACT_EXACT: "The value exactly as the document writes it. Blank if absent.",
    TaskType.EXTRACT_FUZZY: "The value as the document writes it. Blank if absent.",
    TaskType.EXTRACT_LIST: "Every value, separated by semicolons. Blank for none.",
    TaskType.EXTRACT_NUMERIC: "A number with its unit, in any usual form: 1250000 USD, $1.25M. Blank if absent.",
    TaskType.EXTRACT_DATE: "A date in any usual form: 2024-06-15, June 15, 2024, or just 2024-06. Blank if absent.",
    TaskType.SPAN: (
        "Paste the passage, copied from the text file. The harness finds its position. "
        "Blank if the document has no such passage."
    ),
}


class LabelingError(DatasetError):
    """Raised when the labeling plan or sheet is missing, malformed or inconsistent."""


@dataclass
class LabelPlan:
    """Which documents get labeled, and which of them are the holdout.

    Written once by ``sample-labels``, before any model sees a document.
    """

    seed: int
    corpus_size: int
    sampled: list[str]
    holdout: list[str]
    # documents whose sheet rows were filled in with model answers: labeled by correction
    prefilled: list[str] = field(default_factory=list)
    # documents a labeler set aside, with the reason they gave
    skipped: dict[str, str] = field(default_factory=dict)
    created_at: str = ""

    def __post_init__(self) -> None:
        outside = sorted(set(self.holdout) - set(self.sampled))
        if outside:
            raise LabelingError(f"holdout document(s) not in the labeling sample: {', '.join(outside)}")
        shown = sorted(set(self.prefilled) & set(self.holdout))
        if shown:
            raise LabelingError(
                f"holdout document(s) were shown model answers: {', '.join(shown)}; their labels are no longer blind"
            )

    @property
    def inclusion_probability(self) -> float:
        """Return the chance each document had of being drawn, for corpus-level estimates."""
        return len(self.sampled) / self.corpus_size if self.corpus_size else 1.0

    @property
    def to_correct(self) -> list[str]:
        """Return the sampled documents outside the holdout, which may be prefilled."""
        holdout = set(self.holdout)
        return [doc_id for doc_id in self.sampled if doc_id not in holdout]

    def to_dict(self) -> dict[str, Any]:
        """Return the label_plan.json payload."""
        return {
            "seed": self.seed,
            "corpus_size": self.corpus_size,
            "created_at": self.created_at or datetime.now(UTC).isoformat(timespec="seconds"),
            "sampled": list(self.sampled),
            "holdout": list(self.holdout),
            "prefilled": sorted(self.prefilled),
            "skipped": dict(sorted(self.skipped.items())),
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> LabelPlan:
        """Build a plan from the contents of label_plan.json."""
        try:
            return cls(
                seed=int(data["seed"]),
                corpus_size=int(data["corpus_size"]),
                sampled=list(data["sampled"]),
                holdout=list(data["holdout"]),
                prefilled=list(data.get("prefilled", [])),
                skipped=dict(data.get("skipped", {})),
                created_at=str(data.get("created_at", "")),
            )
        except KeyError as exc:
            raise LabelingError(f"label_plan.json is missing key {exc}") from exc


def load_plan(path: Path) -> LabelPlan:
    """Read label_plan.json."""
    if not path.exists():
        raise LabelingError(f"{path} does not exist; run `doc-harness sample-labels` first")
    return LabelPlan.from_dict(json.loads(path.read_text(encoding="utf-8")))


def write_plan(path: Path, plan: LabelPlan) -> None:
    """Write label_plan.json."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(plan.to_dict(), indent=2) + "\n", encoding="utf-8")
    logger.info("wrote %s", path)


def draw_sample(doc_ids: Sequence[str], count: int, holdout_share: float, seed: int) -> LabelPlan:
    """Draw the documents to label, and the holdout among them, at random.

    :param doc_ids: Every document in the corpus
    :param count: How many to label; capped at the corpus size
    :param holdout_share: The fraction of the sample to hold out
    :param seed: The fixed seed, recorded in the plan
    :returns: The plan, with the sample and the holdout in corpus order
    """
    if count <= 0:
        raise LabelingError("the labeling sample must have at least one document")
    if not 0.0 < holdout_share < 1.0:
        raise LabelingError(f"the holdout share must be between 0 and 1, not {holdout_share}")
    corpus = sorted(set(doc_ids))
    if not corpus:
        raise LabelingError("the corpus is empty; run `doc-harness extract` first")
    size = min(count, len(corpus))
    if size < count:
        logger.warning("asked for %d documents but the corpus has %d; labeling all of them", count, len(corpus))
    rng = random.Random(seed)
    sampled = sorted(rng.sample(corpus, size))
    holdout_size = max(1, round(size * holdout_share)) if size > 1 else 0
    holdout = sorted(rng.sample(sampled, holdout_size))
    return LabelPlan(seed=seed, corpus_size=len(corpus), sampled=sampled, holdout=holdout)


def render_cell(task: Any, value: Any, document: str = "") -> str:
    """Render a stored label or a model answer as the text a person would write.

    :param task: The task spec
    :param value: A gold label from labels.jsonl, or a typed model output
    :param document: The document's extracted text, to show a stored span as its passage
    :returns: The cell text; empty for an abstention
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, list | tuple):
        return f"{LIST_SEPARATOR} ".join(render_cell(task, item) for item in value)
    if isinstance(value, Quantity):
        return _render_quantity(value.value, value.unit)
    if isinstance(value, PartialDate):
        return value.value
    if isinstance(value, Span):
        return value.text or ""
    if isinstance(value, Mapping):
        if "value" in value and TaskType(task.type) is TaskType.EXTRACT_NUMERIC:
            return _render_quantity(value["value"], value.get("unit"))
        if "value" in value:
            return str(value["value"])
        if "text" in value and value["text"]:
            return str(value["text"])
        if "start" in value and "end" in value and document:
            return document[int(value["start"]) : int(value["end"])]
    # a model sometimes answers "MSA-1014" with the quote marks; accepted as-is, they would become gold
    return _unquote(str(value))


def _render_quantity(amount: Any, unit: Any) -> str:
    try:
        number = float(amount)
    except (TypeError, ValueError):
        return f"{amount} {unit or ''}".strip()
    # thousands separators, and never scientific notation: 1.25e+06 is not how anyone writes money
    text = f"{int(number):,}" if number.is_integer() else f"{number:,.10f}".rstrip("0")
    return f"{text} {unit}" if unit else text


class CellError(ValueError):
    """Raised when a cell cannot be read as a label for its task."""


def parse_cell(task: Any, text: str, document: str) -> tuple[Any, list[str]]:
    """Read one cell as a gold label in its canonical stored form.

    :param task: The task spec
    :param text: The cell text, already stripped
    :param document: The document's extracted text, for locating spans
    :returns: The label, and any warnings worth showing the labeler
    :raises CellError: If the cell cannot be read as a label for this task
    """
    task_type = TaskType(task.type)
    params = task.params
    warnings: list[str] = []
    blank = not text or (task_type is not TaskType.MULTICLASS and is_null(text, params))
    if task_type in (TaskType.MULTILABEL, TaskType.EXTRACT_LIST):
        pieces = [] if blank else text.replace("\n", LIST_SEPARATOR).split(LIST_SEPARATOR)
        items = [_unquote(item) for item in pieces]
        items = [item for item in items if item]
        if task_type is TaskType.EXTRACT_LIST:
            warnings += _joined_entries(items)
            return items, warnings
        members = []
        for item in items:
            member = normalize_enum(item, params)
            if member not in (task.enum_members or []):
                raise CellError(f"{item!r} is not one of: {', '.join(task.enum_members or [])}")
            members.append(member)
        return sorted(set(members)), warnings

    if task_type is TaskType.MULTICLASS:
        # checked before blankness, so an enum member spelled like a null word survives
        member = normalize_enum(text, params) if text else None
        if member is None:
            return _blank(task), warnings
        if member not in (task.enum_members or []):
            raise CellError(f"{text!r} is not one of: {', '.join(task.enum_members or [])}")
        return member, warnings

    if blank:
        return _blank(task), warnings

    if task_type is TaskType.BINARY:
        answer = boolean(text, params)
        if not isinstance(answer, bool):
            raise CellError(f"{text!r} is not yes or no")
        return answer, warnings
    if task_type in (TaskType.EXTRACT_EXACT, TaskType.EXTRACT_FUZZY):
        return _unquote(text), warnings
    if task_type is TaskType.EXTRACT_NUMERIC:
        quantity = normalize_numeric(text, params)
        if not isinstance(quantity, Quantity):
            raise CellError(f"{text!r} is not a number")
        declared = params.get("unit")
        if declared and quantity.unit and quantity.unit != declared:
            warnings.append(f"unit {quantity.unit!r} differs from the task's declared unit {declared!r}")
        return {"value": quantity.value, "unit": quantity.unit}, warnings
    if task_type is TaskType.EXTRACT_DATE:
        parsed = normalize_date(text, params)
        if not isinstance(parsed, PartialDate):
            raise CellError(f"{text!r} is not a date")
        return parsed.value, warnings
    if task_type is TaskType.SPAN:
        located = locate_quote(text, document)
        if located is None:
            raise CellError(
                "the passage is not in the extracted text; copy it from the text file rather than the PDF, "
                "and paste it without edits"
            )
        # a short passage can match in the wrong place, and span scoring is by position
        if locate_quote(text, document[located.end :]) is not None:
            warnings.append("the passage appears more than once, and the first is used; paste a longer passage")
        return {"start": located.start, "end": located.end}, warnings
    raise CellError(f"no reader for task type {task_type}")


# two entries typed into one cell: "Stryker Corporation and Conformis Inc" is one gold value
# that nothing can ever match, and it caps the task's score without showing up as an error
_JOINED = re.compile(r"\s+and\s+|\s+&\s+|\s*/\s*|,\s+", re.IGNORECASE)


def _looks_like_a_name(part: str) -> bool:
    """Return whether a fragment could stand alone as an entry, rather than being part of one.

    "Johnson & Johnson" and "Procter and Gamble" are single names whose halves are single
    words; "Conformis Inc" and "Jane Q. Smith" stand on their own.
    """
    tokens = part.split()
    if len(tokens) < 2:
        return False
    words = part.casefold().replace(".", "").split()
    if any(word in LEGAL_SUFFIXES for word in words):
        return True
    return sum(1 for token in tokens if token[:1].isupper()) >= 2


def _joined_entries(items: Sequence[str]) -> list[str]:
    """Warn about a cell that looks like several entries written as one."""
    warnings: list[str] = []
    for item in items:
        parts = [part.strip() for part in _JOINED.split(item) if part.strip()]
        if sum(1 for part in parts if _looks_like_a_name(part)) > 1:
            warnings.append(
                f"{item!r} may be more than one entry; separate them with '{LIST_SEPARATOR}' if so. "
                "Two entries in one cell can never be matched, and cap this task's score"
            )
    return warnings


def _unquote(text: str) -> str:
    """Remove quote marks around a whole value, including the curly ones Excel's autocorrect types."""
    stripped = text.strip()
    for opening, closing in (("\u201c", "\u201d"), ("\u2018", "\u2019")):
        if len(stripped) > 2 and stripped.startswith(opening) and stripped.endswith(closing):
            return stripped[1:-1].strip()
    return strip_wrapping_quotes(stripped)


def _blank(task: Any) -> Any:
    """Return the stored abstention for a task, refusing where the task may not abstain."""
    if not task.is_nullable:
        raise CellError("this task may not be blank; tasks.yaml declares it non-nullable")
    return None


def _cell_text(value: Any) -> str:
    """Turn whatever a spreadsheet cell holds into the text the person meant."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class SheetRow:
    """One document's row, as the labeler left it."""

    doc_id: str
    cells: dict[str, str]
    notes: str = ""

    @property
    def skip_reason(self) -> str | None:
        """Return the reason given in a "skip: ..." note, or None if the row is not skipped."""
        match = _SKIP_NOTE.match(self.notes)
        return match.group(1).strip() if match else None

    @property
    def is_empty(self) -> bool:
        """Return whether nothing at all was filled in: not labeled yet, not a considered blank."""
        return not self.notes and not any(self.cells.values())


def write_sheet(
    path: Path,
    registry: Registry,
    plan: LabelPlan,
    values: Mapping[str, Mapping[str, Any]],
    notes: Mapping[str, str],
    text_dir: Path,
    file_names: Mapping[str, str] | None = None,
) -> None:
    """Write the labeling sheet: one tab, one row per sampled document, one column per task.

    :param path: Where to write the .xlsx file
    :param registry: The parsed tasks.yaml
    :param plan: The labeling plan
    :param values: Per document, the answers to show: existing labels, or model answers for
        documents being corrected. Never model answers for the holdout
    :param notes: Per document, notes to carry over
    :param text_dir: The text cache, to show an already-labeled span as its passage
    :param file_names: Per document, its PDF's file name; ``<doc_id>.pdf`` where not given
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    holdout = set(plan.holdout)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME
    task_ids = registry.ids
    headers = [FILE_COLUMN, *task_ids, NOTES_COLUMN]
    sheet.append(headers)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        if header in task_ids:
            task = registry.by_id(header)
            cell.comment = Comment(f"{task.question.strip()}\n\n{FORMAT_HINTS[TaskType(task.type)]}", "doc-harness")
    sheet.cell(row=1, column=1).comment = Comment(
        "Shaded rows start empty on purpose: label them from the document alone, without looking at any "
        "model output. Unshaded rows, if filled in, hold the model's answers: check and correct every cell.\n\n"
        "A blank cell means the document gives no answer.",
        "doc-harness",
    )
    sheet.cell(row=1, column=len(headers)).comment = Comment(
        "Anything worth recording about the document. To set a document aside -- not a contract, "
        "unreadable -- write: skip: <reason>",
        "doc-harness",
    )

    has_spans = any(TaskType(task.type) is TaskType.SPAN for task in registry)
    blind_fill = PatternFill(start_color=BLIND_FILL, end_color=BLIND_FILL, fill_type="solid")
    for doc_id in plan.sampled:
        row_values = values.get(doc_id, {})
        text_file = text_dir / f"{doc_id}.md"
        # an imported span is stored as offsets, and goes back out as the passage it covers
        document = text_file.read_text(encoding="utf-8") if has_spans and text_file.exists() else ""
        sheet.append(
            [
                (file_names or {}).get(doc_id, f"{doc_id}.pdf"),
                *(render_cell(registry.by_id(task_id), row_values.get(task_id), document) for task_id in task_ids),
                notes.get(doc_id, ""),
            ]
        )
        if doc_id in holdout or doc_id not in plan.prefilled:
            sheet.cell(row=sheet.max_row, column=1).fill = blind_fill

    last_row = max(sheet.max_row, 2)
    for cells in sheet.iter_rows(min_row=2, max_row=last_row):
        for cell in cells:
            # text, so Excel leaves leading zeros and dates exactly as typed
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    def add_choices(column_index: int, choices: Sequence[str]) -> None:
        formula = '"' + ",".join(choices) + '"'
        # Excel refuses a list longer than 255 characters; such a column is checked on import instead
        if len(formula) > 255:
            return
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        # a warning, not a stop: a synonym such as "California" is fine and is mapped on import
        validation.errorStyle = "warning"
        validation.error = f"Expected one of: {', '.join(choices)}"
        sheet.add_data_validation(validation)
        letter = get_column_letter(column_index)
        validation.add(f"{letter}2:{letter}{last_row}")

    for offset, task_id in enumerate(task_ids):
        task = registry.by_id(task_id)
        task_type = TaskType(task.type)
        if task_type is TaskType.BINARY:
            add_choices(2 + offset, ["yes", "no"])
        elif task_type is TaskType.MULTICLASS:
            add_choices(2 + offset, task.enum_members or [])

    for column, header in enumerate(headers, start=1):
        is_span = header in task_ids and TaskType(registry.by_id(header).type) is TaskType.SPAN
        width = {FILE_COLUMN: 45, NOTES_COLUMN: 30}.get(header, 50 if is_span else 22)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "B2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    logger.info("wrote %s with %d document(s)", path, len(plan.sampled))


def read_sheet(path: Path, registry: Registry) -> list[SheetRow]:
    """Read the labeler's rows back from an .xlsx workbook or a .csv export of it.

    The first tab is read, whatever it is called, and columns that are not tasks are ignored,
    so a labeler can rename the tab or add a column of their own.

    :param path: The sheet
    :param registry: The parsed tasks.yaml, to check every task has a column
    :returns: One row per document, in sheet order
    """
    if not path.exists():
        raise LabelingError(f"{path} does not exist; run `doc-harness label-sheet` first")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            raw_rows = [dict(row) for row in csv.DictReader(handle)]
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        workbook.close()
        if not rows:
            raise LabelingError(f"{path} is empty")
        names = [_cell_text(value) for value in rows[0]]
        raw_rows = [dict(zip(names, row, strict=False)) for row in rows[1:]]

    headers_seen = set(raw_rows[0]) if raw_rows else set()
    id_column = FILE_COLUMN if FILE_COLUMN in headers_seen else DOC_COLUMN
    missing = [name for name in (id_column, *registry.ids) if name not in headers_seen]
    if raw_rows and missing:
        raise LabelingError(f"{path} is missing column(s): {', '.join(missing)}")
    parsed: list[SheetRow] = []
    for raw in raw_rows:
        name = _cell_text(raw.get(id_column))
        doc_id = doc_id_from_name(name) if id_column == FILE_COLUMN else name
        if not doc_id:
            # a trailing row Excel kept after a deletion
            if any(_cell_text(value) for value in raw.values()):
                raise LabelingError(f"{path}: a row with answers has no {id_column}")
            continue
        parsed.append(
            SheetRow(
                doc_id=doc_id,
                cells={task_id: _cell_text(raw.get(task_id)) for task_id in registry.ids},
                notes=_cell_text(raw.get(NOTES_COLUMN)),
            )
        )
    logger.info("read %d row(s) from %s", len(parsed), path)
    return parsed


@dataclass
class ImportResult:
    """What an import found: the labels it can write, and everything wrong with the sheet."""

    records: list[LabelRecord] = field(default_factory=list)
    skipped: dict[str, str] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        """Return whether the whole sheet was finished and readable."""
        return not self.problems


def import_rows(
    registry: Registry,
    plan: LabelPlan,
    rows: Sequence[SheetRow],
    texts: Mapping[str, str],
) -> ImportResult:
    """Check the whole sheet and turn every row into a label record, collecting every problem.

    The sheet is imported whole or not at all: every sampled document needs a row, and every
    row needs to be finished. A row with nothing in it -- no answer and no note -- is taken
    as not labeled yet, since a document that truly answers none of the questions can say so
    in its note. Problems are gathered rather than raised one at a time, because a person
    fixing a sheet wants the whole list, not one error per run.

    :param registry: The parsed tasks.yaml
    :param plan: The labeling plan the sheet was made from
    :param rows: The rows read back from the sheet
    :param texts: Document id to extracted text, for locating pasted passages
    :returns: The records to write and the problems that block writing them
    """
    result = ImportResult()
    sampled = set(plan.sampled)
    prefilled = set(plan.prefilled)
    seen: set[str] = set()
    for row in rows:
        if row.doc_id in seen:
            result.problems.append(f"{row.doc_id}: appears in more than one row")
            continue
        seen.add(row.doc_id)
        if row.doc_id not in sampled:
            result.problems.append(f"{row.doc_id}: not in the labeling sample in {PLAN_FILE}")
            continue
        reason = row.skip_reason
        if reason is not None:
            if not reason:
                result.problems.append(f"{row.doc_id}: skipped with no reason; write skip: <reason> in {NOTES_COLUMN}")
            result.skipped[row.doc_id] = reason
            continue
        if row.is_empty:
            result.problems.append(
                f"{row.doc_id}: nothing filled in yet. If the document really answers none of the "
                f"questions, say so in {NOTES_COLUMN}"
            )
            continue
        text = texts.get(row.doc_id)
        if text is None:
            result.problems.append(f"{row.doc_id}: no cached text; run `doc-harness extract`")
            continue
        labels: dict[str, Any] = {}
        for task in registry:
            try:
                value, warnings = parse_cell(task, row.cells.get(task.id, ""), text)
            except CellError as exc:
                result.problems.append(f"{row.doc_id} / {task.id}: {exc}")
                continue
            labels[task.id] = value
            result.warnings += [f"{row.doc_id} / {task.id}: {warning}" for warning in warnings]
        if len(labels) == len(registry):
            result.records.append(
                LabelRecord(
                    doc_id=row.doc_id,
                    labels=labels,
                    stratum="random",
                    inclusion_probability=plan.inclusion_probability,
                    # a document whose row showed model answers was labeled by correcting them
                    labeling_mode="corrected" if row.doc_id in prefilled else "blind",
                    notes=row.notes,
                )
            )
    missing = sorted(sampled - seen)
    if missing:
        result.problems.append(f"{len(missing)} sampled document(s) have no row: {', '.join(missing)}")
    return result
